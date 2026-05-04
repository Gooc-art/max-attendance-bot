from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Iterable, List
from urllib.request import Request, urlopen

from app.config import BASE_DIR
from app.integrations.max.client import MaxApiClient
from app.schemas.max import MaxAttachment
from app.services.skud_import import SkudImportResult, SkudImportService


SAFE_FILENAME_RE = re.compile(r"[^A-Za-z0-9._-]+")
NS_MAIN = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
NS_REL = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}


@dataclass
class AttachmentImportResult:
    filename: str
    status: str
    message: str
    import_result: SkudImportResult | None = None


class MaxFileImportService:
    def __init__(self) -> None:
        self.max_client = MaxApiClient.from_settings()
        self.skud_import_service = SkudImportService()
        self.upload_dir = BASE_DIR / "data" / "uploads" / "skud"
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    def import_attachments(self, attachments: Iterable[MaxAttachment], created_by_user_id: str) -> List[AttachmentImportResult]:
        results: List[AttachmentImportResult] = []
        for attachment in attachments:
            if attachment.type != "file":
                continue
            results.append(self._import_file_attachment(attachment, created_by_user_id))
        return results

    def _import_file_attachment(self, attachment: MaxAttachment, created_by_user_id: str) -> AttachmentImportResult:
        if not attachment.payload or not attachment.payload.url:
            return AttachmentImportResult(
                filename=attachment.filename or "unknown",
                status="error",
                message="Во вложении нет URL для скачивания.",
            )

        filename = attachment.filename or "attachment"
        try:
            file_bytes = self._download_file(attachment.payload.url)
            saved_path = self._store_file(filename, file_bytes)
            suffix = Path(filename).suffix.lower()

            if suffix == ".csv":
                result = self.skud_import_service.import_csv(
                    filename=filename,
                    content=file_bytes,
                    created_by_user_id=created_by_user_id,
                )
                return AttachmentImportResult(
                    filename=filename,
                    status="imported",
                    message=self._format_import_message(result),
                    import_result=result,
                )

            if suffix == ".xlsx":
                csv_bytes = self._xlsx_to_csv_bytes(file_bytes)
                result = self.skud_import_service.import_csv(
                    filename=filename,
                    content=csv_bytes,
                    created_by_user_id=created_by_user_id,
                )
                return AttachmentImportResult(
                    filename=filename,
                    status="imported",
                    message=self._format_import_message(result),
                    import_result=result,
                )

            return AttachmentImportResult(
                filename=filename,
                status="unsupported",
                message="Файл `{0}` сохранен в `{1}`, но формат пока не поддержан. Поддерживаются `.csv` и `.xlsx`.".format(
                    filename,
                    saved_path.name,
                ),
            )
        except Exception as exc:
            return AttachmentImportResult(
                filename=filename,
                status="error",
                message="Не удалось обработать файл `{0}`: {1}".format(filename, exc),
            )

    def _download_file(self, url: str) -> bytes:
        headers = {"User-Agent": "attendance-bot/0.1.0"}
        request = Request(url=url, headers=headers, method="GET")
        try:
            with urlopen(request, timeout=60) as response:
                return response.read()
        except Exception:
            if self.max_client.token:
                request = Request(
                    url=url,
                    headers={
                        "User-Agent": "attendance-bot/0.1.0",
                        "Authorization": self.max_client.token,
                    },
                    method="GET",
                )
                with urlopen(request, timeout=60) as response:
                    return response.read()
            raise

    def _store_file(self, filename: str, content: bytes) -> Path:
        source = Path(filename)
        suffix = source.suffix
        stem = source.stem if suffix else source.name
        safe_stem = SAFE_FILENAME_RE.sub("_", stem).strip("._") or "attachment"
        safe_suffix = SAFE_FILENAME_RE.sub("", suffix) if suffix else ""
        safe_name = f"{safe_stem}{safe_suffix}"
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        target = self.upload_dir / f"{timestamp}_{safe_name}"
        target.write_bytes(content)
        return target

    def _xlsx_to_csv_bytes(self, content: bytes) -> bytes:
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                output = StringIO()
                rows = sheet.iter_rows(values_only=True)
                for row in rows:
                    values = ["" if value is None else str(value) for value in row]
                    output.write(";".join(values))
                    output.write("\n")
                return output.getvalue().encode("utf-8")
            finally:
                workbook.close()
        except Exception:
            return self._xlsx_to_csv_bytes_fallback(content)

    def _xlsx_to_csv_bytes_fallback(self, content: bytes) -> bytes:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            shared_strings = self._read_shared_strings(archive)
            sheet_path = self._find_first_sheet_path(archive)
            root = ET.fromstring(archive.read(sheet_path))

        output = StringIO()
        for row in root.findall(".//x:sheetData/x:row", NS_MAIN):
            values: List[str] = []
            expected_index = 0
            for cell in row.findall("x:c", NS_MAIN):
                ref = cell.attrib.get("r", "")
                column_index = self._column_index_from_ref(ref)
                while expected_index < column_index:
                    values.append("")
                    expected_index += 1
                values.append(self._read_cell_value(cell, shared_strings))
                expected_index += 1
            output.write(";".join(values))
            output.write("\n")
        return output.getvalue().encode("utf-8")

    def _read_shared_strings(self, archive: zipfile.ZipFile) -> List[str]:
        if "xl/sharedStrings.xml" not in archive.namelist():
            return []
        root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        items: List[str] = []
        for si in root.findall("x:si", NS_MAIN):
            text_parts = [node.text or "" for node in si.findall(".//x:t", NS_MAIN)]
            items.append("".join(text_parts))
        return items

    def _find_first_sheet_path(self, archive: zipfile.ZipFile) -> str:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        workbook_rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        first_sheet = workbook_root.find("x:sheets/x:sheet", NS_MAIN)
        if first_sheet is None:
            raise ValueError("В XLSX не найден ни один лист")
        rel_id = first_sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        if not rel_id:
            raise ValueError("У листа XLSX отсутствует relationship id")
        target = None
        for rel in workbook_rels.findall("r:Relationship", NS_REL):
            if rel.attrib.get("Id") == rel_id:
                target = rel.attrib.get("Target")
                break
        if not target:
            raise ValueError("Не удалось определить путь к первому листу XLSX")
        return f"xl/{target}".replace("xl/xl/", "xl/")

    def _read_cell_value(self, cell: ET.Element, shared_strings: List[str]) -> str:
        cell_type = cell.attrib.get("t")
        if cell_type == "inlineStr":
            texts = [node.text or "" for node in cell.findall(".//x:t", NS_MAIN)]
            return "".join(texts)
        value = cell.findtext("x:v", default="", namespaces=NS_MAIN)
        if cell_type == "s":
            try:
                return shared_strings[int(value)]
            except Exception:
                return value
        return value

    def _column_index_from_ref(self, ref: str) -> int:
        letters = "".join(ch for ch in ref if ch.isalpha()).upper()
        result = 0
        for ch in letters:
            result = result * 26 + (ord(ch) - ord("A") + 1)
        return max(result - 1, 0)

    @staticmethod
    def _format_import_message(result: SkudImportResult) -> str:
        parts = [
            "Импорт `{filename}` завершен.".format(filename=result.filename),
            "Период: {start} - {end}".format(
                start=result.period_start or "-",
                end=result.period_end or "-",
            ),
            "Уникальных сотрудников: {count}".format(count=result.unique_employees),
            "Строк обработано: {total}".format(total=result.total_rows),
            "Кейсов создано: {created}".format(created=result.created_cases),
            "Ошибок: {errors}".format(errors=result.error_rows),
        ]
        if result.samples:
            parts.append("Примеры:\n" + "\n".join(f"- {item}" for item in result.samples))
        return "\n".join(parts)
