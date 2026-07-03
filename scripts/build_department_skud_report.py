from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.db.base import SessionLocal
from app.db.models import ApprovalModel, AttendanceCaseModel, AttendanceDayRecordModel, EmployeeDirectoryModel, SubstitutionDraftModel
from app.integrations.google_sheets.directory import GoogleSheetsDirectoryService


REQUIRED_HEADERS = {"Сотрудник", "Дата", "Начало дня", "Конец дня"}
CANONICAL_DEPARTMENTS = [
    "0. Руководство",
    "0.1. Зам_руководителя",
    "1. ОКПКиПиБ",
    "2. ОРП",
    "3. ООО",
    "4. ОФиХО",
    "5. ОИТиИБ",
    "6. ОВА",
    "7. ООГУ",
    "8. ОРиУН",
    "9. ОКК",
    "10. ОККСНР",
    "11. ОККНИ",
    "12. ОКННДФЛ",
    "13. ОККНД",
    "14. ОККНДС",
    "15. ОНК№1",
    "16. ОНК№2",
    "17. ОНК№3",
    "18. ООК",
    "19. АО",
    "20. ОУЗ",
    "21. ООПиБ",
    "22. ПО№1",
    "23. ПО№2",
]


@dataclass
class DayRow:
    employee: str
    department: str
    event_date: str
    late_minutes: int
    early_leave_minutes: int
    absence_minutes: int
    first_entry: str
    last_exit: str
    no_show_reason: str
    comment: str
    missing_entry: bool
    missing_exit: bool


@dataclass(frozen=True)
class EmployeeRow:
    employee: str
    department: str


def build_report(input_path: Path, output_path: Path, *, canonical_tabs: bool = False) -> Path:
    rows, period_title = read_skud_rows(input_path)
    department_by_employee = load_department_map()
    rows = [
        DayRow(
            **{
                **row.__dict__,
                "department": department_by_employee.get(row.employee) or fallback_department(row.department),
            }
        )
        for row in rows
    ]

    workbook = Workbook()
    review_sheet = workbook.active
    review_sheet.title = "К проверке"
    write_rows(review_sheet, build_day_review_rows(rows))

    summary_sheet = workbook.create_sheet("Сводка")
    write_rows(summary_sheet, build_department_summary(rows, period_title, canonical_tabs=canonical_tabs))

    for department, department_rows in ordered_department_groups(group_by_department(rows), canonical_tabs=canonical_tabs):
        sheet = workbook.create_sheet(unique_sheet_name(workbook.sheetnames, department))
        write_rows(sheet, build_day_issue_rows(department_rows))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def build_all_time_report(output_path: Path, *, canonical_tabs: bool = False) -> Path:
    with SessionLocal() as session:
        rows = session.query(AttendanceCaseModel).order_by(
            AttendanceCaseModel.event_date.asc(),
            AttendanceCaseModel.employee_name.asc(),
        ).all()
        day_records = session.query(AttendanceDayRecordModel).order_by(
            AttendanceDayRecordModel.event_date.asc(),
            AttendanceDayRecordModel.employee_name.asc(),
        ).all()
        approvals = session.query(ApprovalModel).order_by(ApprovalModel.created_at.asc()).all()
        substitutions = session.query(SubstitutionDraftModel).order_by(SubstitutionDraftModel.created_at.asc()).all()
    employees = load_employee_rows()
    day_records_by_employee = group_day_records_by_employee(day_records)
    approvals_by_case = {approval.case_id: approval for approval in approvals}
    substitutions_by_employee: dict[str, list[SubstitutionDraftModel]] = defaultdict(list)
    for substitution in substitutions:
        substitutions_by_employee[substitution.target_manager].append(substitution)
    employees_by_department = group_employees_by_department(employees, canonical_tabs=canonical_tabs)
    case_groups = group_cases_by_department(rows)
    day_record_groups = group_day_records_by_department(day_records)
    for department in employees_by_department:
        case_groups.setdefault(department, [])
    for department in day_record_groups:
        case_groups.setdefault(department, [])

    workbook = Workbook()
    review_sheet = workbook.active
    review_sheet.title = "К проверке"
    write_rows(review_sheet, build_case_review_rows(rows, approvals_by_case, day_records))

    summary_sheet = workbook.create_sheet("Сводка")
    write_rows(
        summary_sheet,
        build_case_department_summary(
            rows,
            approvals_by_case,
            substitutions,
            employees_by_department=employees_by_department,
            canonical_tabs=canonical_tabs,
        ),
    )

    for department, department_rows in ordered_department_groups(case_groups, canonical_tabs=canonical_tabs):
        sheet = workbook.create_sheet(unique_sheet_name(workbook.sheetnames, department))
        write_rows(
            sheet,
            build_case_issue_rows(
                department_rows,
                approvals_by_case,
                day_records=day_record_groups.get(department, []),
            ),
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def read_skud_rows(input_path: Path) -> tuple[list[DayRow], str]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header_row = find_header_row(sheet.iter_rows(values_only=True))
    if header_row is None:
        raise ValueError("Не найдены заголовки отчета СКУД")

    row_number, headers = header_row
    index = {name: offset for offset, name in enumerate(headers) if name}
    raw_rows: list[DayRow] = []
    for row in sheet.iter_rows(min_row=row_number + 1, values_only=True):
        if not any(cell_text(value) for value in row):
            continue
        employee = cell_text(row[index["Сотрудник"]])
        event_date = cell_text(row[index["Дата"]])
        if not employee or not event_date:
            continue
        department = cell_text(row[index.get("Подразделение", -1)]) if "Подразделение" in index else "Не указано"
        first_entry = cell_text(row[index["Начало дня"]])
        last_exit = cell_text(row[index["Конец дня"]])
        raw_rows.append(
            DayRow(
                employee=employee,
                department=department or "Не указано",
                event_date=event_date,
                late_minutes=duration_minutes(row_value(row, index, "Опоздал")),
                early_leave_minutes=duration_minutes(row_value(row, index, "Ранних уходов")),
                absence_minutes=duration_minutes(row_value(row, index, "Прогулял")),
                first_entry=first_entry,
                last_exit=last_exit,
                no_show_reason=cell_text(row_value(row, index, "Причины не выхода")),
                comment=cell_text(row_value(row, index, "Комм. причины отсутствия")),
                missing_entry=is_missing_entry(first_entry),
                missing_exit=is_missing_exit(last_exit),
            )
        )

    # ponytail: duplicate employee-date rows happen in SKUD exports; keep the row with real marks.
    deduped = {}
    for row in raw_rows:
        key = (row.employee, row.event_date)
        current = deduped.get(key)
        if current is None or row_score(row) > row_score(current):
            deduped[key] = row

    period_title = cell_text(sheet["A1"].value) or "Отчет СКУД"
    return list(deduped.values()), period_title


def find_header_row(rows: Iterable[tuple]) -> tuple[int, list[str]] | None:
    for row_number, row in enumerate(rows, start=1):
        headers = [cell_text(value) for value in row]
        if REQUIRED_HEADERS.issubset(set(headers)):
            return row_number, headers
    return None


def load_department_map() -> dict[str, str]:
    with SessionLocal() as session:
        rows = session.query(EmployeeDirectoryModel.employee_name, EmployeeDirectoryModel.department_name).all()
    return {str(name).strip(): str(department).strip() for name, department in rows if str(name).strip()}


def load_employee_rows() -> list[EmployeeRow]:
    records = GoogleSheetsDirectoryService.from_settings().list_employees(refresh=True)
    if records:
        return [
            EmployeeRow(employee=record.employee_name, department=canonical_department(record.department_name))
            for record in records
            if record.employee_name and getattr(record, "is_active", True)
        ]
    with SessionLocal() as session:
        rows = session.query(EmployeeDirectoryModel.employee_name, EmployeeDirectoryModel.department_name, EmployeeDirectoryModel.is_active).all()
    return [
        EmployeeRow(employee=str(name).strip(), department=canonical_department(str(department or "")))
        for name, department, is_active in rows
        if str(name).strip() and is_active
    ]


def build_department_summary(rows: list[DayRow], title: str, *, canonical_tabs: bool = False) -> list[list[object]]:
    grouped = group_by_department(rows)
    result = [["Параметр", "Значение"], ["Название", title], ["Всего сотрудников", len({row.employee for row in rows})], []]
    result.append(["Отдел", "Сотрудников", "Дней", "Опозданий", "Нет входа", "Нет выхода", "Прогул/отсутствие, мин"])
    for department, department_rows in ordered_department_groups(grouped, canonical_tabs=canonical_tabs):
        result.append(
            [
                department,
                len({row.employee for row in department_rows}),
                len(department_rows),
                sum(1 for row in department_rows if row.late_minutes > 0),
                sum(1 for row in department_rows if row.missing_entry),
                sum(1 for row in department_rows if row.missing_exit),
                sum(row.absence_minutes for row in department_rows),
            ]
        )
    return result


def build_employee_summary(rows: list[DayRow]) -> list[list[object]]:
    by_employee: dict[str, list[DayRow]] = defaultdict(list)
    for row in rows:
        by_employee[row.employee].append(row)

    result = [[
        "Сотрудник",
        "Дней",
        "Опозданий",
        "Опоздал, мин",
        "Ранних уходов, мин",
        "Нет входа",
        "Нет выхода",
        "Прогул/отсутствие, мин",
        "Первый вход",
        "Последний выход",
        "Даты",
        "Причины/комментарии",
    ]]
    for employee, employee_rows in sorted(by_employee.items()):
        employee_rows.sort(key=lambda row: parse_date(row.event_date) or date.min)
        reasons = sorted(
            {
                text
                for row in employee_rows
                for text in (row.no_show_reason, row.comment)
                if text
            }
        )
        result.append(
            [
                employee,
                len(employee_rows),
                sum(1 for row in employee_rows if row.late_minutes > 0),
                sum(row.late_minutes for row in employee_rows),
                sum(row.early_leave_minutes for row in employee_rows),
                sum(1 for row in employee_rows if row.missing_entry),
                sum(1 for row in employee_rows if row.missing_exit),
                sum(row.absence_minutes for row in employee_rows),
                first_clock(row.first_entry for row in employee_rows),
                last_clock(row.last_exit for row in employee_rows),
                ", ".join(row.event_date for row in employee_rows),
                "; ".join(reasons),
            ]
        )
    return result


def build_day_issue_rows(rows: list[DayRow]) -> list[list[object]]:
    result = [issue_header()]
    for row in sorted(rows, key=lambda item: (parse_date(item.event_date) or date.min, item.employee)):
        if row.late_minutes <= 0 and row.early_leave_minutes <= 0:
            continue
        result.append([
            row.event_date,
            row.employee,
            issue_kind(row.late_minutes, row.early_leave_minutes),
            row.first_entry or "-",
            "",
            row.late_minutes or "",
            row.last_exit or "-",
            "",
            row.early_leave_minutes or "",
            "; ".join(value for value in (row.no_show_reason, row.comment) if value),
            "",
        ])
    return result


def build_case_issue_rows(
    rows: list[AttendanceCaseModel],
    approvals_by_case: dict[int, ApprovalModel],
    day_records: list[AttendanceDayRecordModel],
) -> list[list[object]]:
    result = [issue_header()]
    cases_by_key: dict[tuple[str, str], list[AttendanceCaseModel]] = defaultdict(list)
    for row in rows:
        cases_by_key[(row.employee_name or "", row.event_date or "")].append(row)

    seen: set[tuple[str, str]] = set()
    for record in sorted(day_records, key=lambda item: (item.event_date or "", item.employee_name or "")):
        is_late = str(record.arrival_status or "") == "late"
        is_early_leave = str(record.departure_status or "") == "early"
        if not is_late and not is_early_leave:
            continue
        key = (record.employee_name or "", record.event_date or "")
        seen.add(key)
        matched_cases = cases_by_key.get(key, [])
        result.append([
            record.event_date or "-",
            record.employee_name or "-",
            issue_kind(int(record.arrival_delta_minutes or 0) if is_late else 0, abs(int(record.departure_delta_minutes or 0)) if is_early_leave else 0),
            record.first_entry_time or "-",
            record.planned_start or "",
            int(record.arrival_delta_minutes or 0) if is_late else "",
            record.last_exit_time or "-",
            record.planned_end or "",
            abs(int(record.departure_delta_minutes or 0)) if is_early_leave else "",
            join_case_reasons(matched_cases),
            merge_case_approvals(matched_cases, approvals_by_case),
        ])

    for row in sorted(rows, key=lambda item: (item.event_date or "", item.employee_name or "")):
        key = (row.employee_name or "", row.event_date or "")
        if key in seen or row.case_type != "late":
            continue
        result.append([
            row.event_date or "-",
            row.employee_name or "-",
            case_problem(row),
            row.actual_start or "-",
            row.planned_start or "",
            int(row.delay_minutes or 0),
            "-",
            "",
            "",
            row.employee_reason or "",
            approval_label(approval_status(row, approvals_by_case)),
        ])
    return result


def issue_header() -> list[str]:
    return ["Дата", "Сотрудник", "Случай", "Приход", "План прихода", "Опоздание, мин", "Уход", "План ухода", "Ранний уход, мин", "Объяснение", "Согласование"]


def issue_kind(late_minutes: int, early_leave_minutes: int) -> str:
    parts = []
    if late_minutes > 0:
        parts.append("Опоздание")
    if early_leave_minutes > 0:
        parts.append("Ранний уход")
    return "; ".join(parts) or "Случай"


def join_case_reasons(rows: list[AttendanceCaseModel]) -> str:
    return "; ".join(sorted({row.employee_reason for row in rows if row.employee_reason}))


def merge_case_approvals(rows: list[AttendanceCaseModel], approvals_by_case: dict[int, ApprovalModel]) -> str:
    return "; ".join(
        sorted({
            label
            for row in rows
            if (label := approval_label(approval_status(row, approvals_by_case)))
        })
    )


def build_day_review_rows(rows: list[DayRow]) -> list[list[object]]:
    result = [review_header()]
    for row in sorted(rows, key=lambda item: (item.department, item.employee, parse_date(item.event_date) or date.min)):
        problems = day_problems(row)
        if not problems:
            continue
        action = day_action(row)
        result.append([
            priority_label(day_priority(row)),
            action,
            problem_group(problems),
            row.employee,
            row.department,
            row.event_date,
            "-" if row.missing_entry else row.first_entry,
            "-" if row.missing_exit else row.last_exit,
            "; ".join(problems),
            "; ".join(value for value in (row.no_show_reason, row.comment) if value),
            "",
        ])
    return sort_review_rows(result)


def build_case_review_rows(
    rows: list[AttendanceCaseModel],
    approvals_by_case: dict[int, ApprovalModel],
    day_records: list[AttendanceDayRecordModel],
) -> list[list[object]]:
    day_record_by_key = {(record.employee_name or "", record.event_date or ""): record for record in day_records}
    result = [review_header()]
    for row in sorted(rows, key=lambda item: (item.department_name or "", item.employee_name or "", item.event_date or "")):
        approval = approval_status(row, approvals_by_case)
        record = day_record_by_key.get((row.employee_name or "", row.event_date or ""))
        problem = case_problem(row)
        action = case_action(row, approval)
        result.append([
            priority_label(case_priority(row, approval)),
            action,
            problem_group([problem]),
            row.employee_name or "-",
            canonical_department(row.department_name or "Не указано"),
            row.event_date or "-",
            getattr(record, "first_entry_time", "") or getattr(row, "actual_start", "") or "-",
            getattr(record, "last_exit_time", "") or "-",
            problem,
            row.employee_reason or "",
            approval_label(approval),
        ])
    return sort_review_rows(result)


def review_header() -> list[str]:
    return ["Приоритет", "Что сделать", "Группа", "Сотрудник", "Отдел", "Дата", "Приход", "Уход", "Проблема", "Объяснение", "Согласование"]


def day_problems(row: DayRow) -> list[str]:
    problems = []
    if row.missing_entry:
        problems.append("Нет входа")
    if row.missing_exit:
        problems.append("Нет выхода")
    if row.late_minutes > 0:
        problems.append(f"Опоздание {row.late_minutes} мин")
    if row.early_leave_minutes > 0:
        problems.append(f"Ранний уход {row.early_leave_minutes} мин")
    if row.absence_minutes > 0:
        problems.append(f"Отсутствие {row.absence_minutes} мин")
    return problems


def day_action(row: DayRow) -> str:
    if row.missing_entry or row.missing_exit:
        return "Проверить отметки"
    if row.no_show_reason or row.comment:
        return "Проверить причину"
    return "Запросить причину"


def day_priority(row: DayRow) -> int:
    if row.missing_entry or row.missing_exit or row.absence_minutes > 0:
        return 1
    if row.late_minutes > 0 or row.early_leave_minutes > 0:
        return 2
    return 3


def priority_label(priority: int) -> str:
    return {1: "Высокий", 2: "Средний"}.get(priority, "Низкий")


def problem_group(problems: list[str]) -> str:
    text = "; ".join(problems).lower()
    if "нет вход" in text or "нет выход" in text:
        return "Отметки"
    if "отсутств" in text:
        return "Отсутствие"
    if "карточ" in text:
        return "Карточка"
    if "опоздан" in text:
        return "Опоздание"
    if "ранний уход" in text:
        return "Ранний уход"
    return "Другое"


def sort_review_rows(rows: list[list[object]]) -> list[list[object]]:
    header, body = rows[0], rows[1:]
    priority_order = {"Высокий": 0, "Средний": 1, "Низкий": 2}
    body.sort(key=lambda row: (priority_order.get(str(row[0]), 9), str(row[4]), str(row[3]), str(row[5])))
    return [header] + body


def build_case_department_summary(
    rows: list[AttendanceCaseModel],
    approvals_by_case: dict[int, ApprovalModel] | None = None,
    substitutions: list[SubstitutionDraftModel] | None = None,
    employees_by_department: dict[str, list[EmployeeRow]] | None = None,
    *,
    canonical_tabs: bool = False,
) -> list[list[object]]:
    approvals_by_case = approvals_by_case or {}
    substitutions = substitutions or []
    employees_by_department = employees_by_department or {}
    dates = [row.event_date for row in rows if row.event_date]
    result = [
        ["Параметр", "Значение"],
        ["Название", "Отчет СКУД по подразделениям за все время"],
        ["Период с", min(dates) if dates else "-"],
        ["Период по", max(dates) if dates else "-"],
        ["Всего сотрудников", len({row.employee_name for row in rows})],
        [],
        [
            "Подразделение",
            "Сотрудников",
            "Случаев",
            "Опозданий",
            "Минут опоздания",
            "Отсутствий",
            "Отпусков",
            "Командировок",
            "Согласовано",
            "На согласовании",
            "Отклонено",
            "Замещений",
        ],
    ]
    for department, department_rows in ordered_department_groups(group_cases_by_department(rows), canonical_tabs=canonical_tabs):
        employee_names = {row.employee_name for row in department_rows}
        employee_names.update(employee.employee for employee in employees_by_department.get(department, []))
        department_substitutions = [row for row in substitutions if row.target_manager in {case.employee_name for case in department_rows}]
        result.append(
            [
                department,
                len(employee_names),
                len(department_rows),
                sum(1 for row in department_rows if row.case_type == "late"),
                sum(int(row.delay_minutes or 0) for row in department_rows),
                sum(1 for row in department_rows if row.case_type == "absent"),
                sum(1 for row in department_rows if is_reason(row, "отпуск")),
                sum(1 for row in department_rows if is_reason(row, "командиров")),
                sum(1 for row in department_rows if approval_status(row, approvals_by_case) == "approved"),
                sum(1 for row in department_rows if approval_status(row, approvals_by_case) == "pending"),
                sum(1 for row in department_rows if approval_status(row, approvals_by_case) == "rejected"),
                len(department_substitutions),
            ]
        )
    return result


def build_case_employee_summary(
    rows: list[AttendanceCaseModel],
    approvals_by_case: dict[int, ApprovalModel] | None = None,
    substitutions_by_employee: dict[str, list[SubstitutionDraftModel]] | None = None,
    employees: list[EmployeeRow] | None = None,
    day_records_by_employee: dict[str, list[AttendanceDayRecordModel]] | None = None,
) -> list[list[object]]:
    approvals_by_case = approvals_by_case or {}
    substitutions_by_employee = substitutions_by_employee or {}
    day_records_by_employee = day_records_by_employee or {}
    by_employee: dict[str, list[AttendanceCaseModel]] = defaultdict(list)
    for employee in employees or []:
        by_employee[employee.employee]
    for row in rows:
        by_employee[row.employee_name or "-"].append(row)

    result = [[
        "Сотрудник",
        "Всего событий",
        "Опозданий",
        "Минут опоздания",
        "Макс. опоздание, мин",
        "Первый приход",
        "Последний уход",
        "Задержался, раз",
        "Минут задержки",
        "Отсутствий",
        "Отпусков",
        "Командировок",
        "Больничных",
        "Согласовано",
        "На согласовании",
        "Отклонено",
        "Ждет причину",
        "Замещений",
        "Последняя причина",
        "Кто согласовал",
        "Комментарий согласующего",
    ]]
    for employee, employee_rows in sorted(by_employee.items()):
        reasons = sorted({row.employee_reason for row in employee_rows if row.employee_reason})
        approvers = sorted({approval.manager_name for row in employee_rows if (approval := approvals_by_case.get(row.id)) and approval.manager_name})
        comments = sorted({approval.comment for row in employee_rows if (approval := approvals_by_case.get(row.id)) and approval.comment})
        substitutions = substitutions_by_employee.get(employee, [])
        employee_day_records = day_records_by_employee.get(employee, [])
        late_departures = [record for record in employee_day_records if record.departure_status == "late"]
        result.append(
            [
                employee,
                len(employee_rows),
                sum(1 for row in employee_rows if row.case_type == "late"),
                sum(int(row.delay_minutes or 0) for row in employee_rows),
                max([int(row.delay_minutes or 0) for row in employee_rows] or [0]),
                first_entry(employee_day_records),
                last_exit(employee_day_records),
                len(late_departures),
                sum(int(record.departure_delta_minutes or 0) for record in late_departures),
                sum(1 for row in employee_rows if row.case_type == "absent"),
                sum(1 for row in employee_rows if is_reason(row, "отпуск")),
                sum(1 for row in employee_rows if is_reason(row, "командиров")),
                sum(1 for row in employee_rows if is_reason(row, "больнич")),
                sum(1 for row in employee_rows if approval_status(row, approvals_by_case) == "approved"),
                sum(1 for row in employee_rows if approval_status(row, approvals_by_case) == "pending"),
                sum(1 for row in employee_rows if approval_status(row, approvals_by_case) == "rejected"),
                sum(1 for row in employee_rows if row.status == "awaiting_employee_reason"),
                len(substitutions),
                latest_reason(employee_rows, substitutions),
                ", ".join(approvers),
                "; ".join(comments),
            ]
        )
    return result


def approval_status(row: AttendanceCaseModel, approvals_by_case: dict[int, ApprovalModel]) -> str:
    approval = approvals_by_case.get(row.id)
    if approval is not None:
        return approval.status
    if row.status in {"justified", "closed"}:
        return "approved"
    if row.status == "unjustified":
        return "rejected"
    if row.status == "awaiting_manager_approval":
        return "pending"
    return ""


def approval_label(status: str) -> str:
    return {
        "approved": "Согласовано",
        "pending": "На согласовании",
        "rejected": "Отклонено",
    }.get(status, "")


def case_problem(row: AttendanceCaseModel) -> str:
    source = str(getattr(row, "source", "") or "")
    if source == "forgot_chip":
        return "Забыл карточку"
    if source.startswith("late_notice"):
        return "Предупредил об опоздании"
    if row.case_type == "late":
        return f"Опоздание {int(row.delay_minutes or 0)} мин"
    if row.case_type == "absent":
        return "Отсутствие"
    return row.case_type or "Событие"


def case_action(row: AttendanceCaseModel, approval: str) -> str:
    if approval == "pending":
        return "Согласовать или отклонить"
    if approval == "rejected":
        return "Проверить отклонение"
    if not row.employee_reason:
        return "Запросить причину"
    if approval == "approved":
        return "ОК"
    return "Проверить"


def case_priority(row: AttendanceCaseModel, approval: str) -> int:
    if approval in {"pending", "rejected"} or not row.employee_reason:
        return 1
    if row.case_type == "absent" or str(getattr(row, "source", "") or "") == "forgot_chip":
        return 2
    return 3


def is_reason(row: AttendanceCaseModel, pattern: str) -> bool:
    return pattern in (row.employee_reason or "").lower()


def latest_reason(rows: list[AttendanceCaseModel], substitutions: list[SubstitutionDraftModel]) -> str:
    items = [row for row in rows if row.employee_reason]
    if items:
        items.sort(key=lambda row: (row.event_date or "", row.id or 0))
        return items[-1].employee_reason
    if substitutions:
        return substitutions[-1].reason
    return ""


def write_rows(sheet, rows: list[list[object]]) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column[:50])
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 4, 12), 70)


def group_by_department(rows: list[DayRow]) -> dict[str, list[DayRow]]:
    grouped: dict[str, list[DayRow]] = defaultdict(list)
    for row in rows:
        grouped[row.department or "Не указано"].append(row)
    return grouped


def group_cases_by_department(rows: list[AttendanceCaseModel]) -> dict[str, list[AttendanceCaseModel]]:
    grouped: dict[str, list[AttendanceCaseModel]] = defaultdict(list)
    for row in rows:
        grouped[canonical_department(row.department_name or "Не указано")].append(row)
    return grouped


def group_employees_by_department(rows: list[EmployeeRow], *, canonical_tabs: bool = False) -> dict[str, list[EmployeeRow]]:
    grouped: dict[str, list[EmployeeRow]] = defaultdict(list)
    for row in rows:
        department = canonical_department(row.department)
        if canonical_tabs and department not in CANONICAL_DEPARTMENTS:
            department = "Не сопоставлено"
        grouped[department].append(row)
    return grouped


def group_day_records_by_employee(rows: list[AttendanceDayRecordModel]) -> dict[str, list[AttendanceDayRecordModel]]:
    grouped: dict[str, list[AttendanceDayRecordModel]] = defaultdict(list)
    for row in rows:
        grouped[row.employee_name or "-"].append(row)
    return grouped


def group_day_records_by_department(rows: list[AttendanceDayRecordModel]) -> dict[str, list[AttendanceDayRecordModel]]:
    grouped: dict[str, list[AttendanceDayRecordModel]] = defaultdict(list)
    for row in rows:
        grouped[canonical_department(row.department_name or "Не указано")].append(row)
    return grouped


def first_entry(rows: list[AttendanceDayRecordModel]) -> str:
    values = [row.first_entry_time for row in rows if row.first_entry_time]
    return min(values) if values else ""


def last_exit(rows: list[AttendanceDayRecordModel]) -> str:
    values = [row.last_exit_time for row in rows if row.last_exit_time]
    return max(values) if values else ""


def ordered_department_groups(grouped: dict[str, list], *, canonical_tabs: bool) -> list[tuple[str, list]]:
    if not canonical_tabs:
        return sorted(grouped.items())
    result = [(department, grouped.get(department, [])) for department in CANONICAL_DEPARTMENTS]
    unmatched = [
        row
        for department, rows in grouped.items()
        if department not in CANONICAL_DEPARTMENTS
        for row in rows
    ]
    if unmatched:
        result.append(("Не сопоставлено", unmatched))
    return result


def fallback_department(value: str) -> str:
    text = value.strip()
    if not text or looks_like_location(text):
        return "Не указано"
    return text


def canonical_department(value: str) -> str:
    text = str(value or "").strip()
    key = department_key(text)
    aliases = {
        department_key("0. Руководитель"): "0. Руководство",
        department_key("1-ОКПКиПиБ"): "1. ОКПКиПиБ",
        department_key("1-ОКПКиПБ"): "1. ОКПКиПиБ",
        department_key("3. ОО"): "3. ООО",
        department_key("5. ОИБиИТ"): "5. ОИТиИБ",
        department_key("12. ОКК НДФЛ"): "12. ОКННДФЛ",
        department_key("22. ПО1"): "22. ПО№1",
        department_key("23. ПО2"): "23. ПО№2",
    }
    if key in aliases:
        return aliases[key]
    for department in CANONICAL_DEPARTMENTS:
        if key == department_key(department):
            return department
    return text or "Не указано"


def department_key(value: str) -> str:
    return re.sub(r"[^0-9a-zа-я]+", "", str(value or "").lower())


def looks_like_location(value: str) -> bool:
    text = value.lower()
    return any(part in text for part in ("ул.", "д.", "г. ", "город "))


def row_score(row: DayRow) -> int:
    return (0 if row.missing_entry else 100) + (0 if row.missing_exit else 100) - row.absence_minutes


def row_value(row: tuple, index: dict[str, int], name: str):
    position = index.get(name)
    if position is None or position >= len(row):
        return ""
    return row[position]


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return str(value).strip()


def duration_minutes(value) -> int:
    if isinstance(value, time):
        return value.hour * 60 + value.minute + (1 if value.second >= 30 else 0)
    text = cell_text(value)
    if not text or text == "Нет":
        return 0
    match = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
    if not match:
        return 0
    seconds = int(match.group(3) or 0)
    return int(match.group(1)) * 60 + int(match.group(2)) + (1 if seconds >= 30 else 0)


def is_missing_entry(value: str) -> bool:
    return "нет вход" in value.lower() or value in {"0:00", "00:00", "00:00:00"}


def is_missing_exit(value: str) -> bool:
    return "нет выхода" in value.lower() or "23:59" in value


def parse_date(value: str) -> date | None:
    try:
        return datetime.strptime(value, "%d.%m.%Y").date()
    except ValueError:
        return None


def first_clock(values: Iterable[str]) -> str:
    clocks = [value for value in values if value and not is_missing_entry(value)]
    return min(clocks) if clocks else "-"


def last_clock(values: Iterable[str]) -> str:
    clocks = [value for value in values if value and not is_missing_exit(value)]
    return max(clocks) if clocks else "-"


def unique_sheet_name(existing: list[str], raw_name: str) -> str:
    base = re.sub(r"[][\\\\/*?:]", " ", raw_name).strip()[:31] or "Не указано"
    if base not in existing:
        return base
    counter = Counter(existing)
    while True:
        suffix = f" {counter[base] + 1}"
        name = f"{base[:31 - len(suffix)]}{suffix}"
        if name not in existing:
            return name
        counter[base] += 1


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path, nargs="?")
    parser.add_argument("--all-time", action="store_true")
    parser.add_argument("--canonical-tabs", action="store_true")
    parser.add_argument("-o", "--output", type=Path, default=Path("data/reports/skud_departments.xlsx"))
    args = parser.parse_args()
    if args.all_time:
        print(build_all_time_report(args.output, canonical_tabs=args.canonical_tabs))
        return
    if args.input is None:
        parser.error("input is required unless --all-time is used")
    print(build_report(args.input, args.output, canonical_tabs=args.canonical_tabs))


if __name__ == "__main__":
    main()
