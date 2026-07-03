from __future__ import annotations

from types import SimpleNamespace

from openpyxl import Workbook
from openpyxl import load_workbook

from scripts.build_department_skud_report import (
    CANONICAL_DEPARTMENTS,
    EmployeeRow,
    build_case_review_rows,
    build_case_issue_rows,
    build_day_review_rows,
    build_day_issue_rows,
    build_case_employee_summary,
    build_case_department_summary,
    build_report,
    fallback_department,
    group_employees_by_department,
    ordered_department_groups,
    read_skud_rows,
)


def test_location_department_falls_back_to_unknown() -> None:
    assert fallback_department("г. Салехард, ул. Губкина, д.6 «А»") == "Не указано"
    assert fallback_department("Отдел продаж") == "Отдел продаж"


def test_read_skud_rows_keeps_duplicate_with_real_marks(tmp_path) -> None:
    path = tmp_path / "skud.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Отчет о нарушениях: c 25.05.2026 по 25.05.2026"])
    ws.append([])
    ws.append(["Фирма", "Подразделение", "Сотрудник", "Дата", "Опоздал", "Ранних уходов", "Прогулял", "Начало дня", "Конец дня"])
    ws.append(["ФНС", "г. Салехард, ул. Губкина, д.6 «А»", "Иванов Иван", "25.05.2026", "00:00:00", "00:00:00", "07:15:00", "Нет входа (0:00)", "Нет выхода (23:59)"])
    ws.append(["ФНС", "г. Салехард, ул. Губкина, д.6 «А»", "Иванов Иван", "25.05.2026", "00:09:00", "00:00:00", "00:29:00", "08:39:00", "17:52:00"])
    wb.save(path)

    rows, _title = read_skud_rows(path)

    assert len(rows) == 1
    assert rows[0].first_entry == "08:39:00"
    assert rows[0].last_exit == "17:52:00"
    assert rows[0].absence_minutes == 29


def test_build_report_adds_review_sheet_first(tmp_path, monkeypatch) -> None:
    input_path = tmp_path / "skud.xlsx"
    output_path = tmp_path / "out.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Отчет о нарушениях: c 25.05.2026 по 25.05.2026"])
    ws.append([])
    ws.append(["Фирма", "Подразделение", "Сотрудник", "Дата", "Опоздал", "Ранних уходов", "Прогулял", "Начало дня", "Конец дня"])
    ws.append(["ФНС", "Отдел продаж", "Иванов Иван", "25.05.2026", "00:09:00", "00:00:00", "00:00:00", "08:39:00", "17:52:00"])
    wb.save(input_path)
    monkeypatch.setattr("scripts.build_department_skud_report.load_department_map", lambda: {})

    build_report(input_path, output_path)

    result = load_workbook(output_path, read_only=True, data_only=True)
    assert result.sheetnames[:2] == ["К проверке", "Сводка"]
    assert result["К проверке"]["A2"].value == "Средний"
    assert result["К проверке"]["B2"].value == "Запросить причину"
    assert result["К проверке"]["I2"].value == "Опоздание 9 мин"
    assert result["Отдел продаж"]["A1"].value == "Дата"
    assert result["Отдел продаж"]["C2"].value == "Опоздание"
    assert result["Отдел продаж"]["F2"].value == 9


def test_build_case_review_rows_shows_approval_and_action() -> None:
    rows = [
        SimpleNamespace(
            id=1,
            department_name="Отдел продаж",
            employee_name="Иванов Иван",
            event_date="2026-05-04",
            case_type="late",
            delay_minutes=14,
            employee_reason="Пробки",
            status="awaiting_manager_approval",
            source="late_notice_matched",
            actual_start="08:44",
        ),
        SimpleNamespace(
            id=2,
            department_name="Отдел продаж",
            employee_name="Петров Петр",
            event_date="2026-05-04",
            case_type="late",
            delay_minutes=0,
            employee_reason="",
            status="awaiting_employee_reason",
            source="forgot_chip",
            actual_start="08:30",
        ),
    ]
    approvals = {1: SimpleNamespace(status="pending")}
    day_records = [SimpleNamespace(employee_name="Иванов Иван", event_date="2026-05-04", first_entry_time="08:45", last_exit_time="17:40")]

    review = build_case_review_rows(rows, approvals, day_records)

    assert review[0] == ["Приоритет", "Что сделать", "Группа", "Сотрудник", "Отдел", "Дата", "Приход", "Уход", "Проблема", "Объяснение", "Согласование"]
    assert review[1] == ["Высокий", "Согласовать или отклонить", "Опоздание", "Иванов Иван", "Отдел продаж", "2026-05-04", "08:45", "17:40", "Предупредил об опоздании", "Пробки", "На согласовании"]
    assert review[2][8] == "Забыл карточку"
    assert review[2][1] == "Запросить причину"


def test_build_day_review_rows_keeps_only_problem_rows() -> None:
    review = build_day_review_rows([
        SimpleNamespace(
            employee="Иванов Иван",
            department="Отдел продаж",
            event_date="25.05.2026",
            late_minutes=9,
            early_leave_minutes=0,
            absence_minutes=0,
            first_entry="08:39:00",
            last_exit="17:52:00",
            no_show_reason="",
            comment="",
            missing_entry=False,
            missing_exit=False,
        ),
        SimpleNamespace(
            employee="Петров Петр",
            department="Отдел продаж",
            event_date="25.05.2026",
            late_minutes=0,
            early_leave_minutes=0,
            absence_minutes=0,
            first_entry="08:20:00",
            last_exit="17:52:00",
            no_show_reason="",
            comment="",
            missing_entry=False,
            missing_exit=False,
        ),
    ])

    assert len(review) == 2
    assert review[1][3] == "Иванов Иван"
    assert review[1][8] == "Опоздание 9 мин"


def test_build_day_issue_rows_keeps_late_and_early_leave_cases_only() -> None:
    rows = [
        SimpleNamespace(employee="Иванов Иван", event_date="25.05.2026", late_minutes=9, early_leave_minutes=0, first_entry="08:39", last_exit="17:52", no_show_reason="", comment=""),
        SimpleNamespace(employee="Петров Петр", event_date="25.05.2026", late_minutes=0, early_leave_minutes=15, first_entry="08:20", last_exit="17:00", no_show_reason="", comment=""),
        SimpleNamespace(employee="Сидоров Сидор", event_date="25.05.2026", late_minutes=0, early_leave_minutes=0, first_entry="08:20", last_exit="17:52", no_show_reason="", comment=""),
    ]

    result = build_day_issue_rows(rows)

    assert len(result) == 3
    assert result[1][2] == "Опоздание"
    assert result[1][5] == 9
    assert result[2][2] == "Ранний уход"
    assert result[2][8] == 15


def test_build_case_issue_rows_uses_day_records_and_late_case_fallback() -> None:
    cases = [
        SimpleNamespace(id=1, employee_name="Иванов Иван", event_date="2026-05-04", case_type="late", delay_minutes=14, employee_reason="Пробки", actual_start="08:44", planned_start="08:30", source="late_notice_matched", status="awaiting_manager_approval"),
        SimpleNamespace(id=2, employee_name="Петров Петр", event_date="2026-05-05", case_type="late", delay_minutes=7, employee_reason="", actual_start="08:37", planned_start="08:30", source="skud_report", status="awaiting_employee_reason"),
    ]
    approvals = {1: SimpleNamespace(status="pending")}
    day_records = [
        SimpleNamespace(employee_name="Иванов Иван", event_date="2026-05-04", arrival_status="late", arrival_delta_minutes=14, departure_status="early", departure_delta_minutes=-20, first_entry_time="08:44", planned_start="08:30", last_exit_time="16:55", planned_end="17:15"),
    ]

    result = build_case_issue_rows(cases, approvals, day_records)

    assert result[1] == ["2026-05-04", "Иванов Иван", "Опоздание; Ранний уход", "08:44", "08:30", 14, "16:55", "17:15", 20, "Пробки", "На согласовании"]
    assert result[2] == ["2026-05-05", "Петров Петр", "Опоздание 7 мин", "08:37", "08:30", 7, "-", "", "", "", ""]


def test_build_case_department_summary_groups_departments() -> None:
    rows = [
        SimpleNamespace(id=1, department_name="Отдел продаж", employee_name="Иванов Иван", event_date="2026-05-04", case_type="late", delay_minutes=14, employee_reason="", status="awaiting_manager_approval"),
        SimpleNamespace(id=2, department_name="Отдел ИТ", employee_name="Сидоров Сидор", event_date="2026-05-05", case_type="absent", delay_minutes=0, employee_reason="Командировка", status="justified"),
    ]
    approvals = {
        1: SimpleNamespace(status="pending"),
        2: SimpleNamespace(status="approved"),
    }

    summary = build_case_department_summary(rows, approvals)

    assert ["Отдел ИТ", 1, 1, 0, 0, 1, 0, 1, 1, 0, 0, 0] in summary
    assert ["Отдел продаж", 1, 1, 1, 14, 0, 0, 0, 0, 1, 0, 0] in summary


def test_ordered_department_groups_can_emit_canonical_tabs() -> None:
    groups = ordered_department_groups({"Отдел продаж": ["row"], "5. ОИТиИБ": ["it"]}, canonical_tabs=True)

    assert [name for name, _rows in groups[: len(CANONICAL_DEPARTMENTS)]] == CANONICAL_DEPARTMENTS
    assert ("5. ОИТиИБ", ["it"]) in groups
    assert groups[-1] == ("Не сопоставлено", ["row"])


def test_employee_without_cases_stays_in_department_sheet() -> None:
    employees = [EmployeeRow(employee="Пустой Сотрудник", department="5. ОИБиИТ")]
    grouped = group_employees_by_department(employees, canonical_tabs=True)

    rows = build_case_employee_summary([], employees=grouped["5. ОИТиИБ"])

    assert rows[1][0] == "Пустой Сотрудник"
    assert rows[1][1:18] == [0, 0, 0, 0, "", "", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]


def test_employee_summary_uses_day_record_times_and_late_departure() -> None:
    rows = build_case_employee_summary(
        [],
        employees=[EmployeeRow(employee="Иванов Иван", department="5. ОИТиИБ")],
        day_records_by_employee={
            "Иванов Иван": [
                SimpleNamespace(first_entry_time="08:40", last_exit_time="17:30", departure_status="late", departure_delta_minutes=15),
                SimpleNamespace(first_entry_time="08:20", last_exit_time="18:10", departure_status="late", departure_delta_minutes=55),
            ]
        },
    )

    assert rows[1][0] == "Иванов Иван"
    assert rows[1][5] == "08:20"
    assert rows[1][6] == "18:10"
    assert rows[1][7] == 2
    assert rows[1][8] == 70
