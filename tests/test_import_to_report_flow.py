from __future__ import annotations

import tempfile
import warnings
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook

from app.db.base import SessionLocal
from app.db.models import AttendanceCaseModel, AttendanceDayRecordModel, ReportModel, SkudImportModel
from app.services.report_service import ReportService
from app.services.skud_import import SkudImportService
from tests.test_skud_import import _FakeGoogleSheetsScheduleService, _FakeSubstitutionRepository


def _cleanup(marker: str) -> None:
    employee_name = f"Тест Импорт {marker}"
    employee_user_id = f"emp-{marker}"
    with SessionLocal() as session:
        for row in session.query(AttendanceCaseModel).filter(AttendanceCaseModel.employee_user_id == employee_user_id).all():
            session.delete(row)
        for row in session.query(AttendanceDayRecordModel).filter(AttendanceDayRecordModel.employee_user_id == employee_user_id).all():
            session.delete(row)
        for row in session.query(SkudImportModel).filter(SkudImportModel.filename.like(f"%{marker}%")).all():
            session.delete(row)
        for row in session.query(ReportModel).filter(ReportModel.user_id == f"mgr-{marker}").all():
            session.delete(row)
        for row in session.query(AttendanceCaseModel).filter(AttendanceCaseModel.employee_name == employee_name).all():
            session.delete(row)
        session.commit()


def test_skud_imported_late_day_reaches_manager_report_xlsx() -> None:
    marker = uuid4().hex
    employee_name = f"Тест Импорт {marker}"
    employee_user_id = f"emp-{marker}"
    manager_user_id = f"mgr-{marker}"
    report = (
        "Отчет по проходам\n"
        "Ф.И.О;Таб.№;Подразделение;Подразделение;Дата и время;Точка;Событие;Зона\n"
        f"{employee_name};;;Отдел тестирования;04.05.2026 08:47:00;Вход;Проход;Здание\n"
        f"{employee_name};;;Отдел тестирования;04.05.2026 17:20:00;Выход;Проход;Внешний мир\n"
    ).encode("utf-8")

    _cleanup(marker)
    try:
        directory_service = SimpleNamespace(
            get_by_employee_name=lambda _employee_name: SimpleNamespace(
                employee_name=employee_name,
                employee_user_id=employee_user_id,
                manager_name="Тест Руководитель",
                manager_user_id=manager_user_id,
                department_name="Отдел тестирования",
                schedule_name="Пятидневка - Девушки 08:30-12:30 / 14:00-17:15",
                schedule_id="ГР003",
                role="employee",
                is_active=True,
            )
        )

        with patch("app.services.skud_import.GoogleSheetsDirectoryService.from_settings", return_value=directory_service), patch(
            "app.services.skud_import.GoogleSheetsScheduleService.from_settings",
            return_value=_FakeGoogleSheetsScheduleService(),
        ), patch(
            "app.services.skud_import.SubstitutionRepository",
            return_value=_FakeSubstitutionRepository(),
        ):
            import_result = SkudImportService().import_csv(
                filename=f"report-{marker}.txt",
                content=report,
                created_by_user_id="sec-test",
            )

        assert import_result.created_cases == 1
        assert import_result.period_start == "2026-05-04"
        assert import_result.period_end == "2026-05-04"

        service = ReportService()
        with tempfile.TemporaryDirectory() as tmpdir:
            service.reports_dir = Path(tmpdir)
            artifact = service.generate_manager_report(
                user_id=manager_user_id,
                role="manager",
                period="custom",
                date_from="2026-05-04",
                date_to="2026-05-04",
            )
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                workbook = load_workbook(artifact.file_path, read_only=True, data_only=True)
                try:
                    cases_values = "\n".join(
                        " | ".join("" if value is None else str(value) for value in row)
                        for row in workbook["Случаи"].iter_rows(values_only=True)
                    )
                finally:
                    workbook.close()

        assert employee_name in artifact.summary_text
        assert "17" in artifact.summary_text
        assert employee_name in cases_values
        assert "Опоздание" in cases_values
        assert "Нет" in cases_values
    finally:
        _cleanup(marker)
